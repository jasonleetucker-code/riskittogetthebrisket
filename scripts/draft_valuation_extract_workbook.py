#!/usr/bin/env python3
"""Extract golden fixtures from a ``Draft Data`` workbook (Batch A evidence).

Reads an owner ``Draft Data`` workbook (the Google Sheets export that
``scheduled-refresh.yml`` syncs to ``CSVs/Draft Data.xlsx``) and writes:

* ``<out-dir>/<stem>.json`` — every non-empty cell of every sheet, formula
  text verbatim AND the cached value, plus a structured view of the pick
  curve (rows 2-73, every intermediate column, ``M1``, the totals row) and
  the source file's sha256;
* ``<out-dir>/<stem>_inputs.json`` — only the formula INPUTS (column B
  rookie values, the F/G/H historical curves, ``M1``, the budget literal).

The workbook is treated as untrusted data: it is parsed with openpyxl and
nothing in it is executed.  Cached values in a Google Sheets export carry 10
significant digits; the fixture records that so tests compare at that
precision for the real-valued columns (C, E, J) and exactly for the integer
and half-dollar columns (D, I, K, L).

Usage::

    python -I scripts/draft_valuation_extract_workbook.py \
        --workbook PATH --out-dir tests/fixtures/draft_valuation \
        --stem workbook_2026_10_17 [--compare "CSVs/Draft Data.xlsx"]

Exit codes: 0 written; 2 the workbook does not have the expected layout.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

EXTRACTOR_VERSION = "draft_valuation_extract_workbook_v1"
CURVE_SHEET = "Draft Data"
FIRST_ROW, LAST_ROW = 2, 73
CURVE_COLUMNS = "ABCDEFGHIJKL"


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _jsonable(v):
    if v is None or isinstance(v, (bool, int, float, str)):
        return v
    return str(v)


def _cells(path: Path) -> dict[str, dict[str, dict[str, object]]]:
    import openpyxl

    wf = openpyxl.load_workbook(path, data_only=False)
    wv = openpyxl.load_workbook(path, data_only=True)
    out: dict[str, dict[str, dict[str, object]]] = {}
    for ws in wf.worksheets:
        sheet: dict[str, dict[str, object]] = {}
        vs = wv[ws.title]
        for row in ws.iter_rows():
            for c in row:
                cached = vs[c.coordinate].value
                if c.value is None and cached is None:
                    continue
                is_formula = isinstance(c.value, str) and c.value.startswith("=")
                sheet[c.coordinate] = {
                    "f": c.value if is_formula else None,
                    "v": _jsonable(cached if is_formula else c.value),
                }
        out[ws.title] = sheet
    return out


def _curve(cells: dict[str, dict[str, object]]) -> dict[str, object]:
    def v(coord):
        return (cells.get(coord) or {}).get("v")

    rows = []
    for r in range(FIRST_ROW, LAST_ROW + 1):
        row = {"row": r}
        for col in CURVE_COLUMNS:
            row[col] = v(f"{col}{r}")
        rows.append(row)
    headers = {col: v(f"{col}1") for col in CURVE_COLUMNS + "M"}
    totals = {coord: cells.get(coord) for coord in ("D74", "F74", "G74", "H74", "L74")}
    return {
        "headers": headers,
        "protectedRoundingParameterM1": v("M1"),
        "rows": rows,
        "totals": totals,
    }


def _diff(a: dict, b: dict) -> list[dict[str, object]]:
    out = []
    for sheet in sorted(set(a) | set(b)):
        sa, sb = a.get(sheet, {}), b.get(sheet, {})
        for coord in sorted(set(sa) | set(sb)):
            if sa.get(coord) != sb.get(coord):
                out.append({"sheet": sheet, "cell": coord, "a": sa.get(coord), "b": sb.get(coord)})
    return out


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--workbook", required=True, type=Path)
    ap.add_argument("--out-dir", required=True, type=Path)
    ap.add_argument("--stem", required=True)
    ap.add_argument("--compare", type=Path, default=None, help="second workbook to cell-diff")
    ap.add_argument("--extracted-at", default=None, help="ISO timestamp (default: now, UTC)")
    args = ap.parse_args(argv)

    cells = _cells(args.workbook)
    if CURVE_SHEET not in cells:
        print(f"no '{CURVE_SHEET}' sheet", file=sys.stderr)
        return 2
    curve = _curve(cells[CURVE_SHEET])
    if any(row["B"] is None for row in curve["rows"]):
        print("column B is incomplete", file=sys.stderr)
        return 2

    source = {
        "fileName": args.workbook.name,
        "sha256": _sha256(args.workbook),
        "sizeBytes": args.workbook.stat().st_size,
        "sheets": list(cells),
        "producer": "Google Sheets xlsx export (cached formula values stored at 10 significant digits)",
        "extractor": EXTRACTOR_VERSION,
        "extractedAt": args.extracted_at or datetime.now(timezone.utc).isoformat(),
    }
    if args.compare is not None:
        diffs = _diff(cells, _cells(args.compare))
        source["comparedWith"] = {
            "fileName": args.compare.name,
            "sha256": _sha256(args.compare),
            "cellDifferences": len(diffs),
            "differences": diffs[:200],
        }

    fixture = {
        "schema": "draft_valuation_workbook_fixture_v1",
        "source": source,
        "curve": curve,
        "cells": cells,
    }
    rows = curve["rows"]
    inputs = {
        "schema": "draft_valuation_workbook_inputs_v1",
        "source": {"fileName": source["fileName"], "sha256": source["sha256"]},
        "budgetDollars": 1200,
        "protectedRoundingParameterM1": curve["protectedRoundingParameterM1"],
        "picks": [row["A"] for row in rows],
        "rookieValues": {"range": "Draft Data!B2:B73", "values": [row["B"] for row in rows]},
        "historicalCurves": [
            {
                "year": int(curve["headers"][col]),
                "range": f"Draft Data!{col}2:{col}73",
                "values": [row[col] for row in rows],
            }
            for col in "FGH"
        ],
    }
    args.out_dir.mkdir(parents=True, exist_ok=True)
    (args.out_dir / f"{args.stem}.json").write_text(
        json.dumps(fixture, indent=1, sort_keys=False, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    (args.out_dir / f"{args.stem}_inputs.json").write_text(
        json.dumps(inputs, indent=1, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    print(f"wrote {args.stem}.json and {args.stem}_inputs.json ({source['sha256']})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
